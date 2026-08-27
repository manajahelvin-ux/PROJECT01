"""Phases 10-11 - exports et persistance SQLite."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest

from transcribe_ai.core.exceptions import ExportError, StorageError
from transcribe_ai.core.models import (
    AIResult,
    AIUsage,
    JobStatus,
    MediaInfo,
    SourceType,
    TranscriptionJob,
    TranscriptionResult,
    TranscriptSegment,
)
from transcribe_ai.io.database import Database, HistoryRepository, TranscriptionRecord
from transcribe_ai.io.exporters import (
    DocxExporter,
    ExportPayload,
    SrtExporter,
    TxtExporter,
    VttExporter,
    export_transcription,
)
from transcribe_ai.utils.timecode import format_srt_timestamp, format_vtt_timestamp


@pytest.fixture
def payload() -> ExportPayload:
    return ExportPayload(
        text="Bonjour et bienvenue.\n\nAujourd'hui, nous parlons de Python.",
        title="Introduction à Python",
        source="https://www.youtube.com/watch?v=abc",
        duration=125.0,
        language="fr",
        created_at=datetime(2026, 3, 12, 14, 30),
        segments=[
            TranscriptSegment(index=0, start=0.0, end=2.5, text="Bonjour et bienvenue."),
            TranscriptSegment(
                index=1, start=2.5, end=6.0, text="Aujourd'hui, nous parlons de Python."
            ),
        ],
        ai_result=AIResult(
            short_summary="Une introduction a Python.",
            detailed_summary="Presentation detaillee du langage Python.",
            key_points=["Python est simple", "Python est polyvalent"],
            keywords=["python", "programmation"],
            usage=AIUsage(calls=3),
        ),
    )


# --------------------------------------------------------------------------- #
# Timecodes
# --------------------------------------------------------------------------- #
class TestTimecodes:
    @pytest.mark.parametrize(
        ("seconds", "expected"),
        [(0, "00:00:00,000"), (1.5, "00:00:01,500"), (3661.25, "01:01:01,250")],
    )
    def test_format_srt(self, seconds: float, expected: str) -> None:
        assert format_srt_timestamp(seconds) == expected

    def test_format_vtt_utilise_un_point(self) -> None:
        assert format_vtt_timestamp(1.5) == "00:00:01.500"


# --------------------------------------------------------------------------- #
# Exports
# --------------------------------------------------------------------------- #
class TestTxtExport:
    def test_contenu_seul(self, payload: ExportPayload, tmp_path: Path) -> None:
        out = TxtExporter().export(payload, tmp_path / "t.txt")
        content = out.read_text(encoding="utf-8")
        assert "Bonjour et bienvenue." in content
        assert "youtube.com" not in content  # aucune metadonnee en TXT

    def test_creation_du_dossier_parent(self, payload: ExportPayload, tmp_path: Path) -> None:
        out = TxtExporter().export(payload, tmp_path / "sous" / "dossier" / "t.txt")
        assert out.exists()


class TestDocxExport:
    def test_document_professionnel(self, payload: ExportPayload, tmp_path: Path) -> None:
        from docx import Document

        out = DocxExporter().export(payload, tmp_path / "t.docx")
        assert out.exists() and out.stat().st_size > 0
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "TRANSCRIBE AI" in text
        assert "TRANSCRIPTION VIDEO" in text
        assert "Bonjour et bienvenue." in text

    def test_metadonnees_dans_le_tableau(self, payload: ExportPayload, tmp_path: Path) -> None:
        from docx import Document

        out = DocxExporter().export(payload, tmp_path / "t.docx")
        cells = [c.text for t in Document(str(out)).tables for r in t.rows for c in r.cells]
        assert "Introduction à Python" in cells
        assert "00:02:05" in cells
        assert "fr" in cells

    def test_section_ia_presente(self, payload: ExportPayload, tmp_path: Path) -> None:
        from docx import Document

        out = DocxExporter().export(payload, tmp_path / "t.docx")
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "Une introduction a Python." in text
        assert "Python est simple" in text


class TestSubtitleExport:
    def test_structure_srt(self, payload: ExportPayload, tmp_path: Path) -> None:
        content = SrtExporter().export(payload, tmp_path / "t.srt").read_text(encoding="utf-8")
        lines = content.splitlines()
        assert lines[0] == "1"
        assert lines[1] == "00:00:00,000 --> 00:00:02,500"
        assert lines[2] == "Bonjour et bienvenue."
        assert "2\n00:00:02,500 --> 00:00:06,000" in content

    def test_structure_vtt(self, payload: ExportPayload, tmp_path: Path) -> None:
        content = VttExporter().export(payload, tmp_path / "t.vtt").read_text(encoding="utf-8")
        assert content.startswith("WEBVTT")
        assert "00:00:00.000 --> 00:00:02.500" in content

    def test_segments_reconstruits_sans_timestamps(self, tmp_path: Path) -> None:
        """Un texte retravaille par l'IA reste exportable en sous-titres."""
        payload = ExportPayload(text="Phrase une. Phrase deux. Phrase trois.", duration=30.0)
        content = SrtExporter().export(payload, tmp_path / "t.srt").read_text(encoding="utf-8")
        assert content.count("-->") == 3
        assert "00:00:00,000" in content

    def test_sans_contenu_erreur_explicite(self, tmp_path: Path) -> None:
        with pytest.raises(ExportError):
            SrtExporter().export(ExportPayload(text=""), tmp_path / "t.srt")


class TestExportDispatcher:
    @pytest.mark.parametrize("fmt", ["txt", "docx", "srt", "vtt"])
    def test_les_quatre_formats(self, payload: ExportPayload, tmp_path: Path, fmt: str) -> None:
        out = export_transcription(payload, tmp_path / f"export.{fmt}")
        assert out.exists() and out.suffix == f".{fmt}"

    def test_format_inconnu(self, payload: ExportPayload, tmp_path: Path) -> None:
        with pytest.raises(ExportError, match="pdf"):
            export_transcription(payload, tmp_path / "export.pdf")

    def test_construction_depuis_un_job(self) -> None:
        job = TranscriptionJob(source="https://x/v", source_type=SourceType.URL)
        job.media_info = MediaInfo(
            title="Ma video", source="https://x/v", source_type=SourceType.URL, duration=60
        )
        job.transcription = TranscriptionResult(text="texte brut", language="fr")
        job.ai_result = AIResult(cleaned_text="Texte nettoye.")
        payload = ExportPayload.from_job(job)
        assert payload.text == "Texte nettoye."  # la version IA prime
        assert payload.title == "Ma video"
        assert payload.default_filename("txt") == "ma-video.txt"


# --------------------------------------------------------------------------- #
# SQLite
# --------------------------------------------------------------------------- #
@pytest.fixture
def repo(tmp_path: Path) -> HistoryRepository:
    return HistoryRepository(Database(f"sqlite:///{tmp_path / 'test.sqlite3'}"))


class TestHistoryRepository:
    def test_schema_complet(self) -> None:
        colonnes = set(TranscriptionRecord.__table__.columns.keys())
        assert {
            "id", "title", "url", "duration", "language", "created_at",
            "word_count", "status", "transcription_path", "audio_path",
        } <= colonnes

    def test_ajout_et_lecture(self, repo: HistoryRepository) -> None:
        record = repo.add(TranscriptionRecord(title="Test", url="https://x", word_count=10))
        assert record.id is not None
        assert repo.get(record.id).title == "Test"

    def test_liste_triee_du_plus_recent(self, repo: HistoryRepository) -> None:
        repo.add(TranscriptionRecord(title="Ancien", created_at=datetime(2020, 1, 1)))
        repo.add(TranscriptionRecord(title="Recent", created_at=datetime(2026, 1, 1)))
        assert [r.title for r in repo.list()] == ["Recent", "Ancien"]

    def test_recherche(self, repo: HistoryRepository) -> None:
        repo.add(TranscriptionRecord(title="Cours de Python", text="variables"))
        repo.add(TranscriptionRecord(title="Recette de cuisine", text="oeufs"))
        assert len(repo.list(search="python")) == 1
        assert len(repo.list(search="oeufs")) == 1

    def test_mise_a_jour_du_texte(self, repo: HistoryRepository) -> None:
        record = repo.add(TranscriptionRecord(title="T", text="ancien", word_count=1))
        repo.update_text(record.id, "un nouveau texte plus long")
        updated = repo.get(record.id)
        assert updated.text == "un nouveau texte plus long"
        assert updated.word_count == 5

    def test_mise_a_jour_dun_id_inexistant(self, repo: HistoryRepository) -> None:
        with pytest.raises(StorageError):
            repo.update_text(9999, "texte")

    def test_suppression(self, repo: HistoryRepository) -> None:
        record = repo.add(TranscriptionRecord(title="A supprimer"))
        assert repo.delete(record.id) is True
        assert repo.get(record.id) is None
        assert repo.delete(record.id) is False

    def test_statistiques(self, repo: HistoryRepository) -> None:
        repo.add(TranscriptionRecord(title="A", word_count=100, duration=60, ai_calls=2))
        repo.add(TranscriptionRecord(title="B", word_count=50, duration=30, ai_calls=3))
        stats = repo.stats()
        assert stats == {"count": 2, "words": 150, "duration": 90.0, "ai_calls": 5}

    def test_enregistrement_dun_job_complet(self, repo: HistoryRepository) -> None:
        job = TranscriptionJob(source="https://x/v", source_type=SourceType.URL)
        job.media_info = MediaInfo(
            title="Video", source="https://x/v", source_type=SourceType.URL,
            duration=120, author="Chaine",
        )
        job.transcription = TranscriptionResult(
            text="texte", language="fr",
            segments=[TranscriptSegment(index=0, start=0, end=1, text="texte")],
        )
        job.ai_result = AIResult(
            cleaned_text="Texte.", short_summary="Resume", key_points=["A", "B"],
            usage=AIUsage(calls=4),
        )
        job.status = JobStatus.COMPLETED

        record = repo.save_job(job)
        assert record.id == job.record_id
        assert record.title == "Video" and record.author == "Chaine"
        assert record.text == "Texte." and record.ai_calls == 4
        assert record.key_points() == ["A", "B"]
        assert len(record.segments()) == 1

    def test_persistance_entre_deux_sessions(self, tmp_path: Path) -> None:
        url = f"sqlite:///{tmp_path / 'persist.sqlite3'}"
        first = HistoryRepository(Database(url))
        first.add(TranscriptionRecord(title="Persistant"))
        first.db.dispose()

        second = HistoryRepository(Database(url))
        assert [r.title for r in second.list()] == ["Persistant"]
