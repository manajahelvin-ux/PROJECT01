"""Phase 10 - exports TXT, DOCX, SRT et VTT.

Toutes les fonctions d'export partagent le meme objet d'entree
(`ExportPayload`), ce qui garantit des metadonnees coherentes entre les
formats et permet d'ajouter un format sans toucher a l'appelant.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from transcribe_ai.core.exceptions import ExportError
from transcribe_ai.core.interfaces import Exporter
from transcribe_ai.core.models import AIResult, MediaInfo, TranscriptSegment
from transcribe_ai.utils.logging_config import get_logger
from transcribe_ai.utils.text import slugify, word_count
from transcribe_ai.utils.timecode import (
    format_duration,
    format_srt_timestamp,
    format_vtt_timestamp,
)

logger = get_logger(__name__)

#: Duree minimale d'un sous-titre : evite les flashs illisibles.
MIN_SUBTITLE_DURATION = 0.5


@dataclass(slots=True)
class ExportPayload:
    """Donnees necessaires a tous les formats d'export."""

    text: str
    title: str = "Transcription"
    source: str = ""
    duration: float = 0.0
    language: str = "auto"
    created_at: datetime = field(default_factory=datetime.now)
    segments: list[TranscriptSegment] = field(default_factory=list)
    ai_result: AIResult | None = None

    @classmethod
    def from_job(cls, job) -> ExportPayload:
        """Construit la charge utile depuis un `TranscriptionJob`."""
        info: MediaInfo | None = job.media_info
        transcription = job.transcription
        ai = job.ai_result
        text = (ai.cleaned_text if ai and ai.cleaned_text else "") or (
            transcription.text if transcription else ""
        )
        return cls(
            text=text,
            title=(info.title if info else job.title) or "Transcription",
            source=job.source,
            duration=(info.duration if info else 0.0)
            or (transcription.duration if transcription else 0.0),
            language=(transcription.language if transcription else job.language) or "auto",
            segments=list(transcription.segments) if transcription else [],
            ai_result=ai,
        )

    @property
    def word_count(self) -> int:
        return word_count(self.text)

    def default_filename(self, extension: str) -> str:
        return f"{slugify(self.title)}.{extension.lstrip('.')}"


def _prepare(destination: Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    return destination


# --------------------------------------------------------------------------- #
# TXT
# --------------------------------------------------------------------------- #
class TxtExporter(Exporter):
    """Texte brut, sans metadonnee : le contenu et rien d'autre."""

    extension = "txt"
    label = "Texte (.txt)"

    def export(self, job_data: ExportPayload, destination: Path) -> Path:
        destination = _prepare(destination)
        try:
            destination.write_text(job_data.text.strip() + "\n", encoding="utf-8")
        except OSError as exc:
            raise ExportError(str(exc), user_message="Ecriture du fichier TXT impossible.") from exc
        return destination


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
class DocxExporter(Exporter):
    """Document Word professionnel : page de garde, metadonnees, contenu."""

    extension = "docx"
    label = "Document Word (.docx)"

    def export(self, job_data: ExportPayload, destination: Path) -> Path:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt, RGBColor
        except ImportError as exc:  # pragma: no cover - dependance declaree
            raise ExportError(
                "python-docx manquant",
                user_message="python-docx n'est pas installe. Executez : pip install python-docx",
            ) from exc

        destination = _prepare(destination)
        document = Document()

        header = document.add_paragraph()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header.add_run("TRANSCRIBE AI")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x6C, 0x5C, 0xE7)

        title = document.add_heading("TRANSCRIPTION VIDEO", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        table = document.add_table(rows=0, cols=2)
        table.style = "Light List Accent 1"
        for label, value in (
            ("Titre", job_data.title),
            ("Source", job_data.source or "-"),
            ("Duree", format_duration(job_data.duration)),
            ("Langue", job_data.language),
            ("Mots", str(job_data.word_count)),
            ("Date", job_data.created_at.strftime("%d/%m/%Y %H:%M")),
        ):
            row = table.add_row().cells
            row[0].paragraphs[0].add_run(label).bold = True
            row[1].text = str(value)

        ai = job_data.ai_result
        if ai and (ai.short_summary or ai.key_points):
            document.add_page_break()
            if ai.short_summary:
                document.add_heading("RESUME", level=1)
                document.add_paragraph(ai.short_summary)
            if ai.detailed_summary:
                document.add_heading("RESUME DETAILLE", level=1)
                document.add_paragraph(ai.detailed_summary)
            if ai.key_points:
                document.add_heading("POINTS CLES", level=1)
                for point in ai.key_points:
                    document.add_paragraph(point, style="List Bullet")
            if ai.keywords:
                document.add_heading("MOTS-CLES", level=1)
                document.add_paragraph(", ".join(ai.keywords))

        document.add_page_break()
        document.add_heading("TRANSCRIPTION", level=1)
        for paragraph in (job_data.text or "").split("\n\n"):
            cleaned = paragraph.strip()
            if cleaned:
                document.add_paragraph(cleaned)

        try:
            document.save(str(destination))
        except OSError as exc:
            raise ExportError(
                str(exc),
                user_message="Ecriture du document Word impossible (fichier deja ouvert ?).",
            ) from exc
        return destination


# --------------------------------------------------------------------------- #
# Sous-titres
# --------------------------------------------------------------------------- #
def _segments_or_fallback(payload: ExportPayload) -> list[TranscriptSegment]:
    """Retourne les segments horodates, ou en fabrique depuis le texte.

    Un texte retravaille par l'IA ne porte plus de timestamps : on repartit
    alors les phrases sur la duree connue, pour rester exportable en SRT/VTT.
    """
    if payload.segments:
        return payload.segments
    from transcribe_ai.utils.text import split_sentences

    sentences = split_sentences(payload.text)
    if not sentences:
        return []
    total = payload.duration or float(len(sentences) * 3)
    total_chars = sum(len(s) for s in sentences) or 1
    segments: list[TranscriptSegment] = []
    cursor = 0.0
    for index, sentence in enumerate(sentences):
        share = len(sentence) / total_chars * total
        end = cursor + max(share, MIN_SUBTITLE_DURATION)
        segments.append(TranscriptSegment(index=index, start=cursor, end=end, text=sentence))
        cursor = end
    return segments


class SrtExporter(Exporter):
    """Sous-titres SubRip."""

    extension = "srt"
    label = "Sous-titres (.srt)"

    def export(self, job_data: ExportPayload, destination: Path) -> Path:
        destination = _prepare(destination)
        segments = _segments_or_fallback(job_data)
        if not segments:
            raise ExportError(
                "Aucun segment", user_message="Aucun contenu horodate a exporter en SRT."
            )
        blocks = []
        for number, seg in enumerate(segments, start=1):
            end = max(seg.end, seg.start + MIN_SUBTITLE_DURATION)
            blocks.append(
                f"{number}\n"
                f"{format_srt_timestamp(seg.start)} --> {format_srt_timestamp(end)}\n"
                f"{seg.text.strip()}\n"
            )
        destination.write_text("\n".join(blocks), encoding="utf-8")
        return destination


class VttExporter(Exporter):
    """Sous-titres WebVTT."""

    extension = "vtt"
    label = "Sous-titres web (.vtt)"

    def export(self, job_data: ExportPayload, destination: Path) -> Path:
        destination = _prepare(destination)
        segments = _segments_or_fallback(job_data)
        if not segments:
            raise ExportError(
                "Aucun segment", user_message="Aucun contenu horodate a exporter en VTT."
            )
        lines = ["WEBVTT", "", f"NOTE Genere par TRANSCRIBE AI - {job_data.title}", ""]
        for number, seg in enumerate(segments, start=1):
            end = max(seg.end, seg.start + MIN_SUBTITLE_DURATION)
            lines.append(str(number))
            lines.append(f"{format_vtt_timestamp(seg.start)} --> {format_vtt_timestamp(end)}")
            lines.append(seg.text.strip())
            lines.append("")
        destination.write_text("\n".join(lines), encoding="utf-8")
        return destination


#: Registre des formats : alimente directement les boutons de l'interface.
EXPORTERS: dict[str, Exporter] = {
    "txt": TxtExporter(),
    "docx": DocxExporter(),
    "srt": SrtExporter(),
    "vtt": VttExporter(),
}


def export_transcription(payload: ExportPayload, destination: Path, fmt: str | None = None) -> Path:
    """Exporte selon le format demande, ou deduit de l'extension du fichier."""
    key = (fmt or Path(destination).suffix.lstrip(".")).lower()
    exporter = EXPORTERS.get(key)
    if exporter is None:
        raise ExportError(
            f"Format inconnu : {key}",
            user_message=f"Format « {key} » non supporte. Disponibles : {', '.join(EXPORTERS)}",
        )
    path = exporter.export(payload, Path(destination))
    logger.info("Export %s : %s", key.upper(), path.name)
    return path
